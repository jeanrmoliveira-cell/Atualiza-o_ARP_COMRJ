"""
Extração de TODAS as Atas de Registro de Preços VIGENTES - COMRJ (UASG 771300)
API Dados Abertos Compras.gov.br (v2) + vigência autoritativa do PNCP.

Descobre dinamicamente TODAS as atas ainda vigentes na data de referência,
inclusive de pregões de anos anteriores cuja ata continua válida E ATAS
PRORROGADAS (renovadas).

Passos:
  1. modulo-arp/2_consultarARPItem  -> todos os itens de ARP da UASG
     (varrendo dataVigenciaInicial dos últimos anos).
  2. VIGÊNCIA AUTORITATIVA: a API de dados abertos NÃO reflete prorrogações
     (mantém a vigência original). Por isso, para cada compra consultamos o
     PNCP (/api/pncp/v1/.../atas), que traz a vigência real e o flag de
     cancelamento, e sobrescrevemos a data de vigência final por ata
     (chave numeroControlePncpAta == numeroControlePNCP do PNCP).
  3. Mantém apenas VIGENTES: vigência final (PNCP) >= HOJE, não excluído,
     não cancelado. Itens prorrogados recebem situação "Vigente (prorrogada)".
  4. Cruza cada compra pela CHAVE CANÔNICA do PNCP
     (item.numeroControlePncpCompra == contratacao.numeroControlePNCP),
     obtendo NUP (processo), objeto, situação, modalidade e valores.
  5. Gera planilha no mesmo layout da versão anterior.
"""

import requests, time
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, date, timedelta

UASG = "771300"
BASE = "https://dadosabertos.compras.gov.br"
PNCP = "https://pncp.gov.br/api/consulta/v1"
PNCP_ATAS = "https://pncp.gov.br/api/pncp/v1"   # vigência autoritativa (reflete prorrogação)
HOJE = date.today()   # data de referência = hoje (edite p/ uma data fixa se quiser)
ANOS_VIGENCIA = list(range(HOJE.year - 3, HOJE.year + 1))  # cobre atas vigentes hoje, inclusive prorrogadas
# Só verifica no PNCP compras cuja vigência (dados abertos) terminou nos últimos
# ~400 dias ou ainda está aberta — janela suficiente p/ captar prorrogações.
LIMIAR_PRORROGA = HOJE - timedelta(days=400)

SESSAO = requests.Session()
SESSAO.headers.update({"Accept": "application/json"})

# ── Categoria: lookup manual (sessão anterior) + heurística por palavra-chave ──
CAT_LOOKUP = {
    ("90044", "2025"): "Saúde", ("90194", "2024"): "Saúde", ("90194", "2025"): "Material Comum",
    ("90064", "2025"): "Gêneros", ("90131", "2025"): "Gêneros",
    ("95006", "2026"): "Gêneros", ("96021", "2025"): "Sobressalente",
    ("90092", "2025"): "Sobressalente", ("90110", "2025"): "Sobressalente",
    ("90135", "2025"): "CLG", ("90103", "2025"): "CLG",
    ("90074", "2024"): "Saúde", ("90049", "2025"): "Saúde",
    ("90079", "2025"): "Saúde", ("90058", "2025"): "Saúde",
    ("90033", "2025"): "Saúde", ("90140", "2025"): "Material Comum",
    ("90119", "2025"): "Material Comum", ("90108", "2025"): "Material Comum",
    ("90148", "2025"): "Material Comum", ("90166", "2025"): "Material Comum",
    ("90008", "2026"): "Fardamento", ("90010", "2026"): "Fardamento",
    ("90182", "2025"): "Viatura", ("90183", "2025"): "Material Comum",
    ("96016", "2025"): "Munição", ("96013", "2025"): "Munição",
    ("96003", "2026"): "Material Comum", ("90050", "2025"): "Material Comum",
    ("90051", "2026"): "Material Comum", ("90080", "2025"): "Saúde",
    ("90050", "2026"): "Gêneros", ("90117", "2025"): "Material Comum",
    ("90015", "2024"): "CLG", ("90035", "2024"): "CLG",
    ("90179", "2025"): "Saúde", ("90056", "2025"): "Saúde",
    ("90056", "2024"): "Saúde", ("90104", "2025"): "Saúde",
    ("90105", "2025"): "CLG", ("90012", "2025"): "Fardamento",
    ("90046", "2026"): "Material Comum", ("90085", "2026"): "Material Comum",
    ("90045", "2026"): "Material Comum", ("90094", "2026"): "Fardamento",
    ("90098", "2026"): "Material Comum", ("90038", "2026"): "Gêneros",
    ("90161", "2025"): "CLG", ("90100", "2026"): "Material Comum",
    ("90090", "2024"): "Material Comum", ("90095", "2025"): "Saúde",
    ("90129", "2025"): "Saúde", ("90171", "2025"): "Sobressalente",
}
CAT_KEYWORDS = [
    ("Saúde",         ["medicament", "cirurg", "odontolog", "hospital", "saude", "saúde",
                        "traumato", "ortope", "hemodin", "neurocir", "opme", "medico",
                        "médico", "laborator", "farmac", "protese", "prótese", "rms",
                        "primeiros socorros", "kpsi", "kit de primeiro"]),
    ("Gêneros",       ["genero", "gênero", "aliment", "hortifr", "frigorif", "racope",
                        "café", "cafe", "carne", "rac e ran", "racao", "ração", "frigorific"]),
    ("Fardamento",    ["fardamento", "uniforme", "calçado", "calcado", "confecç", "confecc",
                        "vestuario", "vestuário", "metais", "insignia", "insígnia"]),
    ("Munição",       ["muniç", "munic", "pirotec", "menos letal", "explosiv"]),
    ("CLG",           ["combustiv", "combustív", "lubrific", "graxa", "qav", "gasolina",
                        "diesel", "óleo diesel", "oleo diesel", "odm", "odr"]),
    ("Viatura",       ["viatura", "automov", "automóv", "veicul", "veícul"]),
    ("Sobressalente", ["sobressalent", "pdr", "reparo", "manutenç", "peças", "peca de reposi",
                        "espia", "riachuelo", "humaita", "humaitá"]),
    ("Material Comum",["tinta", "roupa de cama", "químic", "quimic", "expediente", "higiene",
                        "descart", "material comum", "pintura", "bandeira"]),
]

def categoria(num, ano, objeto):
    c = CAT_LOOKUP.get((str(num), str(ano)))
    if c:
        return c
    o = (objeto or "").lower()
    for cat, kws in CAT_KEYWORDS:
        if any(k in o for k in kws):
            return cat
    return "—"

# ── HTTP helper ───────────────────────────────────────────────
def get_json(url, tentativas=4, timeout=60):
    for t in range(tentativas):
        try:
            r = SESSAO.get(url, timeout=timeout)
            if r.status_code == 200 and r.text.strip():
                try:
                    return r.json()
                except Exception:
                    time.sleep(1.2 * (t + 1)); continue
            if r.status_code == 400:
                return None
            if r.status_code == 404:
                return None
            time.sleep(1.2 * (t + 1))
        except Exception:
            time.sleep(1.5 * (t + 1))
    return None

# ── 1) Itens de ARP (todos, sem filtrar vigência ainda) ───────
def buscar_itens_arp():
    print(f"Buscando itens de ARP (vig. inicial {ANOS_VIGENCIA[0]}-{ANOS_VIGENCIA[-1]})...", flush=True)
    brutos = []
    for y in ANOS_VIGENCIA:
        pagina = 1
        while True:
            url = (f"{BASE}/modulo-arp/2_consultarARPItem"
                   f"?pagina={pagina}&tamanhoPagina=500"
                   f"&codigoUnidadeGerenciadora={UASG}"
                   f"&dataVigenciaInicialMin={y}-01-01"
                   f"&dataVigenciaInicialMax={y}-12-31")
            j = get_json(url, timeout=90)
            if not j or not j.get("resultado"):
                break
            brutos.extend(j["resultado"])
            tp = j.get("totalPaginas", 1)
            print(f"  vig.inicial {y}: pág {pagina}/{tp} (+{len(j['resultado'])})", flush=True)
            if pagina >= tp:
                break
            pagina += 1
            time.sleep(0.1)
        time.sleep(0.12)

    # Uma linha por (ata, item): mantém o fornecedor VENCEDOR (menor
    # classificacaoFornecedor). Demais fornecedores = cadastro reserva (descartado).
    def rank(it):
        try:
            return int(str(it.get("classificacaoFornecedor") or "999"))
        except Exception:
            return 999
    melhor = {}
    for it in brutos:
        k = (it.get("numeroControlePncpAta", ""), it.get("numeroItem", ""))
        cur = melhor.get(k)
        if cur is None or rank(it) < rank(cur):
            melhor[k] = it
    unicos = list(melhor.values())
    print(f"  brutos {len(brutos)} | únicos (ata,item) {len(unicos)}", flush=True)
    return unicos

def _parse_dt(s):
    try:
        return datetime.strptime(str(s)[:10], "%Y-%m-%d").date()
    except Exception:
        return None

# ── 2) Vigência AUTORITATIVA do PNCP (reflete prorrogações) ────
def _compra_para_url(ncp_compra):
    """'00394502000144-1-014311/2025' -> (cnpj, ano, seq)."""
    try:
        p = ncp_compra.split("-")
        cnpj = p[0]
        seq, ano = p[2].split("/")
        return cnpj, ano, int(seq)
    except Exception:
        return None

def carregar_vigencias_pncp(unicos):
    """Para cada compra com ata potencialmente vigente/prorrogada, consulta o
    PNCP e mapeia numeroControlePNCP(ata) -> {vigFim, vigIni, cancelado}."""
    # compras candidatas: alguma ata com vig. final (dados abertos) >= limiar
    compras = {}
    for it in unicos:
        vf = _parse_dt(it.get("dataVigenciaFinal"))
        if vf and vf >= LIMIAR_PRORROGA:
            compras.setdefault(it.get("numeroControlePncpCompra"), True)
    compras = [c for c in compras if c]
    print(f"Consultando vigência autoritativa no PNCP para {len(compras)} compras...", flush=True)
    vig = {}
    falhas = []
    for i, ncp in enumerate(compras, 1):
        parsed = _compra_para_url(ncp)
        if not parsed:
            continue
        cnpj, ano, seq = parsed
        url = f"{PNCP_ATAS}/orgaos/{cnpj}/compras/{ano}/{seq}/atas"
        j = None
        for t in range(3):
            try:
                r = SESSAO.get(url, timeout=30)
                if r.status_code == 200 and r.text.strip():
                    j = r.json(); break
                if r.status_code in (204, 404):
                    j = []; break
                time.sleep(1.0 * (t + 1))
            except Exception:
                time.sleep(1.5 * (t + 1))
        if j is None:
            falhas.append(ncp); continue
        arr = j if isinstance(j, list) else (j.get("data") or [])
        for a in arr:
            k = a.get("numeroControlePNCP")
            if k:
                vig[k] = {
                    "vigFim": _parse_dt(a.get("dataVigenciaFim")),
                    "vigIni": _parse_dt(a.get("dataVigenciaInicio")),
                    "cancelado": bool(a.get("cancelado")),
                }
        if i % 25 == 0:
            print(f"  ... {i}/{len(compras)} compras", flush=True)
        time.sleep(0.08)
    print(f"  atas com vigência PNCP: {len(vig)} | falhas de consulta: {len(falhas)}", flush=True)
    return vig, falhas

def filtrar_vigentes(unicos, vig_pncp):
    """Aplica a vigência autoritativa e mantém só os itens vigentes hoje.
    Marca 'prorrogada' quando o PNCP estende além da data de dados abertos."""
    vigentes = []
    n_prorrog = 0
    for it in unicos:
        if it.get("itemExcluido"):
            continue
        da_fim = _parse_dt(it.get("dataVigenciaFinal"))
        info = vig_pncp.get(it.get("numeroControlePncpAta"))
        prorrogada = False
        if info:
            if info["cancelado"]:
                continue
            vfim = info["vigFim"] or da_fim
            if info["vigFim"] and da_fim and info["vigFim"] > da_fim:
                prorrogada = True
            vini = info["vigIni"] or _parse_dt(it.get("dataVigenciaInicial"))
        else:
            vfim = da_fim
            vini = _parse_dt(it.get("dataVigenciaInicial"))
        if not vfim or vfim < HOJE:
            continue
        # injeta a vigência autoritativa de volta no item (strings ISO)
        it["_vig_fim"] = vfim.strftime("%Y-%m-%d")
        it["_vig_ini"] = vini.strftime("%Y-%m-%d") if vini else (it.get("dataVigenciaInicial") or "—")[:10]
        it["_prorrogada"] = prorrogada
        if prorrogada:
            n_prorrog += 1
        vigentes.append(it)
    print(f"  VIGENTES {len(vigentes)} (das quais prorrogadas: {n_prorrog})", flush=True)
    return vigentes

# ── 2) Pré-carga de contratações (chave canônica numeroControlePNCP) ──
def precarregar_contratacoes(modalidades, anos):
    idx_pncp, idx_numano = {}, {}
    print("Pré-carregando contratações...", flush=True)
    for ano in anos:
        for mod in sorted(modalidades):
            pagina = 1
            while True:
                url = (f"{BASE}/modulo-contratacoes/1_consultarContratacoes_PNCP_14133"
                       f"?pagina={pagina}&tamanhoPagina=500"
                       f"&unidadeOrgaoCodigoUnidade={UASG}"
                       f"&dataPublicacaoPncpInicial={ano}-01-01"
                       f"&dataPublicacaoPncpFinal={ano}-12-31"
                       f"&codigoModalidade={mod}")
                j = get_json(url, timeout=60)
                if not j or not j.get("resultado"):
                    break
                for c in j["resultado"]:
                    idx_pncp[c.get("numeroControlePNCP")] = c
                    idx_numano[(str(c.get("numeroCompra")), str(c.get("anoCompraPncp")))] = c
                if pagina >= j.get("totalPaginas", 1):
                    break
                pagina += 1
                time.sleep(0.1)
            time.sleep(0.1)
    print(f"  {len(idx_pncp)} contratações carregadas.", flush=True)
    return idx_pncp, idx_numano

# ── Fallback: API de Consulta do PNCP por id ──────────────────
def consulta_pncp_por_id(ncp):
    """ncp ex.: '00394502000144-1-009237/2024' -> consulta a compra no PNCP."""
    try:
        parts = ncp.split("-")
        cnpj = parts[0]
        seq, ano = parts[2].split("/")
        seq = int(seq)
    except Exception:
        return None
    url = f"{PNCP}/orgaos/{cnpj}/compras/{ano}/{seq}"
    for t in range(2):
        try:
            r = SESSAO.get(url, timeout=20)
            if r.status_code == 200 and r.text.strip():
                j = r.json()
                # normaliza nomes de campo p/ os usados no módulo de contratações
                return {
                    "objetoCompra": j.get("objetoCompra"),
                    "processo": j.get("processo"),
                    "valorTotalEstimado": j.get("valorTotalEstimado"),
                    "valorTotalHomologado": j.get("valorTotalHomologado"),
                    "modalidadeNome": j.get("modalidadeNome"),
                    "existeResultado": j.get("existeResultado"),
                    "situacaoCompraNomePncp": j.get("situacaoCompraNome") or j.get("situacaoCompraNomePncp"),
                    "_fonte": "PNCP",
                }
            if r.status_code in (400, 404):
                return None
            time.sleep(1.5 * (t + 1))
        except Exception:
            time.sleep(2 * (t + 1))
    return None

# ── Formatações ───────────────────────────────────────────────
def fmt_nup(processo):
    if not processo:
        return "—"
    s = "".join(ch for ch in str(processo) if ch.isdigit())
    if len(s) == 17:
        return f"{s[:5]}.{s[5:11]}/{s[11:15]}-{s[15:17]}"
    return str(processo)

def fmt_cnpj(c):
    if not c:
        return ""
    s = "".join(ch for ch in str(c) if ch.isdigit())
    if len(s) == 14:
        return f"{s[:2]}.{s[2:5]}.{s[5:8]}/{s[8:12]}-{s[12:]}"
    if len(s) == 11:
        return f"{s[:3]}.{s[3:6]}.{s[6:9]}-{s[9:]}"
    return str(c)

def num_or_dash(v):
    try:
        return float(v)
    except Exception:
        return "—"

def link_ata_pncp(it):
    ncp = it.get("numeroControlePncpAta") or ""
    try:
        parts = ncp.split("-")
        cnpj = parts[0]
        seqcomp, ano = parts[2].split("/")
        seqata = parts[3]
        return f"https://pncp.gov.br/app/atas/{cnpj}/{ano}/{int(seqcomp)}/{int(seqata)}"
    except Exception:
        return ""

def situacao_pregao(compra):
    if not compra:
        # há ARP vigente -> a compra foi homologada, mesmo sem registro na listagem
        return {"situacao": "Homologado (ata vigente)*", "estimado": None,
                "homologado": None, "modalidade": ""}
    existe = compra.get("existeResultado")
    homol = compra.get("valorTotalHomologado")
    if existe and homol:
        sit = "Homologado (resultado publicado)"
    elif existe:
        sit = "Publicado - resultado parcial"
    else:
        sit = compra.get("situacaoCompraNomePncp") or "Publicado - aguardando resultado"
    return {"situacao": sit,
            "estimado": compra.get("valorTotalEstimado"),
            "homologado": homol,
            "modalidade": compra.get("modalidadeNome", "")}

# ════════════════════════════════════════════════════════════
print("=" * 70, flush=True)
print("ATAS VIGENTES - COMRJ UASG 771300 | ref:", HOJE.strftime("%d/%m/%Y"), flush=True)
print("=" * 70, flush=True)

unicos = buscar_itens_arp()
vig_pncp, falhas_pncp = carregar_vigencias_pncp(unicos)
itens = filtrar_vigentes(unicos, vig_pncp)

mods, anos_compra = set(), set()
for it in itens:
    try:
        mods.add(int(it.get("codigoModalidadeCompra")))
    except Exception:
        pass
    a = str(it.get("anoCompra"))
    if a.isdigit():
        anos_compra.add(int(a))
mods |= {5, 6, 7}
y0 = min(anos_compra) if anos_compra else 2024
idx_pncp, idx_numano = precarregar_contratacoes(mods, range(y0, 2028))

# Agrupar por compra (chave canônica)
grupos = defaultdict(list)
for it in itens:
    grupos[it.get("numeroControlePncpCompra")].append(it)
print(f"\nCompras (pregões) com atas vigentes: {len(grupos)}", flush=True)

# Resolver a contratação de cada grupo (com fallback PNCP)
compra_cache = {}
def resolver_compra(ncp_compra, sample):
    if ncp_compra in compra_cache:
        return compra_cache[ncp_compra]
    c = idx_pncp.get(ncp_compra)
    if c is None:
        num = str(sample.get("numeroCompra")); ano = str(sample.get("anoCompra"))
        c = idx_numano.get((num, ano))
    if c is None and ncp_compra:
        c = consulta_pncp_por_id(ncp_compra)
        if c:
            print(f"  [PNCP fallback] {ncp_compra} recuperado", flush=True)
    compra_cache[ncp_compra] = c
    return c

# Montar linhas
linhas = []
for ncp_compra, its in grupos.items():
    sample = its[0]
    compra = resolver_compra(ncp_compra, sample)
    num = str((compra or {}).get("numeroCompra") or sample.get("numeroCompra") or "")
    ano = str((compra or {}).get("anoCompraPncp") or sample.get("anoCompra") or "")
    numero_pregao = f"{num}/{ano}" if num and ano else (num or "—")
    sit = situacao_pregao(compra)
    objeto = (compra or {}).get("objetoCompra") or "—"
    nup = fmt_nup((compra or {}).get("processo"))
    cat = categoria(num, ano, objeto)

    def keyit(it):
        try:
            return int("".join(c for c in str(it.get("numeroItem", "")) if c.isdigit()))
        except Exception:
            return 0
    for it in sorted(its, key=keyit):
        qh = num_or_dash(it.get("quantidadeHomologadaItem"))
        qe = num_or_dash(it.get("quantidadeEmpenhada"))
        saldo = (qh - qe) if isinstance(qh, float) and isinstance(qe, float) else "—"
        linhas.append({
            "nup": nup, "numero": numero_pregao, "categoria": cat, "objeto": objeto,
            "sit_pregao": sit["situacao"],
            "modalidade": it.get("nomeModalidadeCompra") or sit["modalidade"] or "—",
            "estimado": num_or_dash(sit["estimado"]), "homologado": num_or_dash(sit["homologado"]),
            "num_item": it.get("numeroItem", "—"), "desc": it.get("descricaoItem", "—"),
            "tipo": it.get("tipoItem", "—"),
            "sit_item": "Vigente (prorrogada)" if it.get("_prorrogada") else "Vigente",
            "qh": qh, "qe": qe, "saldo": saldo,
            "forn": it.get("nomeRazaoSocialFornecedor", "—"), "cnpj": fmt_cnpj(it.get("niFornecedor", "")),
            "vu": num_or_dash(it.get("valorUnitario")), "vt": num_or_dash(it.get("valorTotal")),
            "arp": it.get("numeroAtaRegistroPreco", "—"),
            "vig_ini": it.get("_vig_ini") or (it.get("dataVigenciaInicial") or "—")[:10],
            "vig_fim": it.get("_vig_fim") or (it.get("dataVigenciaFinal") or "—")[:10],
            "link": link_ata_pncp(it),
            "_ano": ano,
        })

# ordenar linhas por ano/numero/item para leitura
linhas.sort(key=lambda d: (d["_ano"], d["numero"]))
print(f"Total de linhas de itens vigentes: {len(linhas)}", flush=True)

# ── Excel ────────────────────────────────────────────────────
print("Gerando Excel...", flush=True)
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Atas Vigentes"

HDR = "003366"; VERDE = "D9EAD3"; AMAR = "FFF2CC"; CINZA = "EFEFEF"; AZUL = "CFE2F3"
bold = Font(bold=True, color="FFFFFF", size=10)
fillh = PatternFill("solid", fgColor=HDR)
thin = Side(style="thin", color="CCCCCC")
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
lft = Alignment(horizontal="left", vertical="center", wrap_text=True)

COLS = [("NUP", 18), ("Nº Pregão", 11), ("Categoria", 13), ("Objeto", 40),
        ("Situação do Pregão", 24), ("Modalidade", 16),
        ("Valor Estimado (R$)", 16), ("Valor Homologado (R$)", 17),
        ("Nº Item", 8), ("Descrição do Item", 50), ("Tipo", 9),
        ("Situação do Item", 13), ("Qtd Homolog.", 12), ("Qtd Empenh.", 12), ("Saldo", 11),
        ("Fornecedor", 36), ("CNPJ", 19), ("Valor Unit. (R$)", 14), ("Valor Total (R$)", 15),
        ("Nº ARP", 11), ("Vig. Início", 11), ("Vig. Fim", 11), ("Link PNCP", 46), ("Extraído em", 16)]

for c, (t, w) in enumerate(COLS, 1):
    cell = ws.cell(1, c, t); cell.font = bold; cell.fill = fillh
    cell.alignment = ctr; cell.border = bd
    ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[1].height = 32
ws.freeze_panes = "A2"

agora = datetime.now().strftime("%d/%m/%Y %H:%M")
CENTER_COLS = {1, 2, 3, 5, 6, 9, 11, 12, 13, 14, 15, 17, 20, 21, 22, 24}
MONEY_COLS = {7, 8, 18, 19}

def cor_sit_pregao(s):
    if "Homologado" in s:
        return VERDE
    if "aguardando" in s or "parcial" in s or "Divulgada" in s:
        return AMAR
    return CINZA

r = 2
for d in linhas:
    valores = [d["nup"], d["numero"], d["categoria"], d["objeto"], d["sit_pregao"],
               d["modalidade"], d["estimado"], d["homologado"], d["num_item"], d["desc"],
               d["tipo"], d["sit_item"], d["qh"], d["qe"], d["saldo"], d["forn"], d["cnpj"],
               d["vu"], d["vt"], d["arp"], d["vig_ini"], d["vig_fim"], d["link"], agora]
    for c, v in enumerate(valores, 1):
        cell = ws.cell(r, c, v); cell.border = bd
        cell.alignment = ctr if c in CENTER_COLS else lft
        if c == 5:
            cell.fill = PatternFill("solid", fgColor=cor_sit_pregao(d["sit_pregao"]))
        elif c == 12:
            cell.fill = PatternFill("solid", fgColor=AZUL if "prorrogada" in str(d["sit_item"]) else VERDE)
        if c in MONEY_COLS and isinstance(v, float):
            cell.number_format = 'R$ #,##0.00'
    r += 1

# ── Aba Resumo por Pregão ────────────────────────────────────
wr = wb.create_sheet("Resumo por Pregão")
RC = [("NUP", 18), ("Nº Pregão", 12), ("Categoria", 13), ("Objeto", 42),
      ("Situação do Pregão", 24), ("Itens Vigentes", 13),
      ("Valor Estimado", 16), ("Valor Homologado", 16), ("Soma Itens Vig. (R$)", 18)]
for c, (t, w) in enumerate(RC, 1):
    cell = wr.cell(1, c, t); cell.font = bold; cell.fill = fillh
    cell.alignment = ctr; cell.border = bd
    wr.column_dimensions[get_column_letter(c)].width = w
wr.freeze_panes = "A2"; wr.row_dimensions[1].height = 28

resumo = {}
for d in linhas:
    k = d["numero"]
    if k not in resumo:
        resumo[k] = {"nup": d["nup"], "cat": d["categoria"], "obj": d["objeto"],
                     "sit": d["sit_pregao"], "est": d["estimado"], "hom": d["homologado"],
                     "n": 0, "soma": 0.0, "_ano": d["_ano"]}
    resumo[k]["n"] += 1
    if isinstance(d["vt"], float):
        resumo[k]["soma"] += d["vt"]
rr = 2
for k, v in sorted(resumo.items(), key=lambda x: (x[1]["_ano"], x[0])):
    row = [v["nup"], k, v["cat"], v["obj"], v["sit"], v["n"], v["est"], v["hom"], v["soma"]]
    for c, val in enumerate(row, 1):
        cell = wr.cell(rr, c, val); cell.border = bd
        cell.alignment = ctr if c != 4 else lft
        if c == 5:
            cell.fill = PatternFill("solid", fgColor=cor_sit_pregao(v["sit"]))
        if c in (7, 8, 9) and isinstance(val, float):
            cell.number_format = 'R$ #,##0.00'
    rr += 1

# ── Aba Legenda ──────────────────────────────────────────────
wl = wb.create_sheet("Legenda")
leg = [("CONCEITO", "SIGNIFICADO"),
       ("Critério de seleção", "Atas com Vig. Fim >= data de referência, item NÃO excluído e ata NÃO cancelada"),
       ("Data de referência", HOJE.strftime("%d/%m/%Y")),
       ("UASG", f"{UASG} - Centro de Obtenção da Marinha no Rio de Janeiro (COMRJ)"),
       ("Fonte", "API Dados Abertos Compras.gov.br (ARP e Contratações) + PNCP (vigência autoritativa)"),
       ("Vigência (Vig. Fim)", "Vem do PNCP, que REFLETE PRORROGAÇÕES. A API de dados abertos guarda só a vigência original; por isso a vigência final é confirmada/atualizada no PNCP por ata."),
       ("", ""),
       ("Situação do Item", ""),
       ("  Vigente", "Ata com vigência em curso nesta data"),
       ("  Vigente (prorrogada)", "Ata cuja vigência foi PRORROGADA/renovada no PNCP além da data original de dados abertos"),
       ("Situação do Pregão", ""),
       ("  Homologado (resultado publicado)", "Resultado homologado registrado no PNCP"),
       ("  Homologado (ata vigente)*", "Há ARP vigente, mas a compra não consta na listagem de contratações (objeto/NUP/valores não recuperados)"),
       ("Saldo", "Qtd Homologada - Qtd Empenhada"),
       ("NUP", "Número Único de Protocolo (campo 'processo' da contratação)"),
       ("Categoria", "Manual (processos conhecidos) ou inferida por palavra-chave do objeto"),
       ("Link PNCP", "Link direto para a Ata no Portal Nacional de Contratações Públicas")]
for i, (a, b) in enumerate(leg, 1):
    wl.cell(i, 1, a).font = Font(bold=(i == 1 or b == ""))
    wl.cell(i, 2, b)
wl.column_dimensions["A"].width = 34
wl.column_dimensions["B"].width = 78

out = r"C:\Users\Jean Macedo\Downloads\COMRJ_Atas_Vigentes.xlsx"
wb.save(out)

# ── Estatísticas ─────────────────────────────────────────────
tot_val = sum(d["vt"] for d in linhas if isinstance(d["vt"], float))
sem_meta = sorted({d["numero"] for d in linhas if "ata vigente)*" in d["sit_pregao"]})
n_prorrog = sum(1 for d in linhas if "prorrogada" in str(d["sit_item"]))
atas_prorrog = sorted({d["arp"] for d in linhas if "prorrogada" in str(d["sit_item"])})
print("\n" + "=" * 70, flush=True)
print(f"OK Salvo: {out}", flush=True)
print(f"   Pregões com atas vigentes: {len(resumo)}", flush=True)
print(f"   Linhas de itens vigentes : {len(linhas)}", flush=True)
print(f"   Itens em atas PRORROGADAS : {n_prorrog}  (atas: {len(atas_prorrog)})", flush=True)
print(f"   Valor total (itens vig.) : R$ {tot_val:,.2f}", flush=True)
if atas_prorrog:
    print(f"   Atas prorrogadas: {', '.join(atas_prorrog[:30])}{' ...' if len(atas_prorrog)>30 else ''}", flush=True)
if falhas_pncp:
    print(f"   [aviso] {len(falhas_pncp)} compras sem confirmação de vigência no PNCP (usada a data de dados abertos)", flush=True)
if sem_meta:
    print(f"   Pregões sem metadados (objeto/NUP): {len(sem_meta)} -> {sem_meta}", flush=True)
print("=" * 70, flush=True)
